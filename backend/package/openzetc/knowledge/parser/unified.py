"""Unified parser module for markdown conversion."""

from __future__ import annotations

import asyncio
import base64
import os
import re
import shutil
import subprocess
import tempfile
import threading
import time
from dataclasses import dataclass, field
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Any

import aiofiles
from docling.datamodel.base_models import InputFormat
from docling.document_converter import DocumentConverter
from langchain_community.document_loaders import PyPDFLoader
from markdownify import markdownify as md_convert

from openzetc.knowledge.parser.zip_utils import process_zip_file as _process_zip_file
from openzetc.storage.minio import get_minio_client
from openzetc.utils import logger

SUPPORTED_FILE_EXTENSIONS: tuple[str, ...] = (
    ".txt",
    ".md",
    ".doc",
    ".docx",
    ".html",
    ".htm",
    ".mhtml",
    ".mht",
    ".json",
    ".csv",
    ".xls",
    ".xlsx",
    ".pdf",
    ".pptx",
    ".jpg",
    ".jpeg",
    ".png",
    ".bmp",
    ".tiff",
    ".tif",
    ".zip",
)


def is_supported_file_extension(file_name: str | os.PathLike[str]) -> bool:
    """Check whether the given file path has a supported extension."""
    return Path(file_name).suffix.lower() in SUPPORTED_FILE_EXTENSIONS


@dataclass(slots=True)
class MarkdownParseResult:
    """统一的 Markdown 解析结果。"""

    markdown: str
    file_ext: str | None = None
    artifacts: dict[str, Any] = field(default_factory=dict)


_docling_converter: DocumentConverter | None = None
_docling_converter_lock = threading.Lock()


def _get_docling_converter() -> DocumentConverter:
    """获取 Docling 文档转换器单例。"""
    global _docling_converter
    if _docling_converter is None:
        _docling_converter = DocumentConverter(
            format_options={
                InputFormat.DOCX: None,
                InputFormat.XLSX: None,
                InputFormat.PPTX: None,
            }
        )
    return _docling_converter


def _resolve_image_storage_params(params: dict | None) -> tuple[str, str]:
    params = params or {}

    image_bucket = params.get("image_bucket") or "public"
    image_prefix = params.get("image_prefix")
    if image_prefix:
        normalized_prefix = str(image_prefix).strip("/")
        if normalized_prefix:
            return image_bucket, normalized_prefix

    return image_bucket, "unknown/kb-images"


def _resolve_ocr_engine_params(params: dict | None) -> tuple[str, dict[str, Any]]:
    from openzetc import config

    params = params or {}
    engine = str(params.get("ocr_engine") if "ocr_engine" in params else config.default_ocr_engine)
    engine = engine.strip() or config.default_ocr_engine
    engine_config = params.get("ocr_engine_config")
    processor_params = dict(params)
    if isinstance(engine_config, dict):
        processor_params.update(engine_config)
    return engine, processor_params


def _upload_image_to_minio(image_data: bytes, filename: str, bucket_name: str, object_prefix: str) -> str:
    """上传图片到 MinIO，返回 URL。"""
    minio_client = get_minio_client()
    minio_client.ensure_bucket_exists(bucket_name)

    normalized_prefix = object_prefix.strip("/") or "unknown/kb-images"
    timestamp = int(time.time() * 1000000)
    object_name = f"{normalized_prefix}/{timestamp}_{Path(filename).name}"

    result = minio_client.upload_file(
        bucket_name=bucket_name,
        object_name=object_name,
        data=image_data,
    )
    return result.url


def _parse_data_uri(data_uri: str) -> tuple[bytes, str]:
    """解析 data URI，返回 (image_data, mime_type)。"""
    header, base64_data = data_uri.split(",", 1)
    mime_type = header.split(":")[1].split(";")[0]
    image_data = base64.b64decode(base64_data)
    return image_data, mime_type


def _convert_with_docling(file_path: Path, params: dict | None = None) -> str:
    """使用 Docling 将 docx/xlsx/pptx 转换为 Markdown。"""
    params = params or {}
    image_bucket, image_prefix = _resolve_image_storage_params(params)

    with _docling_converter_lock:
        converter = _get_docling_converter()
        result = converter.convert(file_path)

    if result.status.name != "SUCCESS":
        raise RuntimeError(f"Docling 转换失败: {result.status}")

    doc = result.document

    if hasattr(doc, "pictures") and doc.pictures:
        replacements: list[str] = []
        for pic in doc.pictures:
            uri = str(pic.image.uri) if hasattr(pic, "image") and hasattr(pic.image, "uri") else ""
            if uri.startswith("data:"):
                filename = "image"
                try:
                    image_data, mime_type = _parse_data_uri(uri)
                    filename = f"image_{int(time.time() * 1000000)}.{mime_type.split('/')[-1]}"
                    url = _upload_image_to_minio(image_data, filename, image_bucket, image_prefix)
                    replacements.append(f"![{filename}]({url})")
                except Exception as e:  # noqa: BLE001
                    logger.error(f"上传图片失败 {filename}: {e}")
                    replacements.append(f"[图片: {filename}]")
            else:
                replacements.append("")

        markdown = doc.export_to_markdown()
        for replacement in replacements:
            markdown = re.sub(r"<!--\s*image\s*-->", replacement, markdown, count=1)
        return markdown

    return doc.export_to_markdown()


def _convert_docx_with_python_docx(file_path: Path) -> str:
    """使用 python-docx 解析 DOCX（Docling 失败时兜底）。"""
    from docx import Document

    document = Document(str(file_path))
    blocks: list[str] = []

    for para in document.paragraphs:
        text = para.text.strip()
        if text:
            blocks.append(text)

    for table in document.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(cells)

        if not rows:
            continue

        header = rows[0]
        blocks.append(f"| {' | '.join(header)} |")
        blocks.append(f"| {' | '.join(['---'] * len(header))} |")

        for row in rows[1:]:
            normalized_row = row + [""] * (len(header) - len(row))
            blocks.append(f"| {' | '.join(normalized_row[: len(header)])} |")

        blocks.append("")

    return "\n\n".join(blocks).strip()


def _convert_doc_with_libreoffice(file_path: Path) -> str:
    """使用 LibreOffice 将旧版 DOC 转为 DOCX 后提取正文。"""
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise RuntimeError("DOC 解析依赖 LibreOffice，请先安装 soffice/libreoffice")

    timeout_seconds = 60
    with tempfile.TemporaryDirectory(prefix="openzetc-doc-") as temp_dir:
        temp_path = Path(temp_dir)
        input_path = temp_path / "source.doc"
        output_path = temp_path / "source.docx"
        profile_path = temp_path / "lo-profile"
        profile_path.mkdir(parents=True, exist_ok=True)
        input_path.write_bytes(file_path.read_bytes())

        command = [
            executable,
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            "--nodefault",
            "--nolockcheck",
            f"-env:UserInstallation={profile_path.resolve().as_uri()}",
            "--convert-to",
            "docx",
            "--outdir",
            str(temp_path),
            str(input_path),
        ]
        try:
            result = subprocess.run(
                command,
                capture_output=True,
                timeout=timeout_seconds,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError(f"DOC 转换超时（{timeout_seconds} 秒）") from exc

        if result.returncode != 0:
            detail = (result.stderr or result.stdout).decode("utf-8", errors="ignore").strip()
            raise RuntimeError(f"DOC 转换失败: {detail or 'LibreOffice 执行失败'}")
        if not output_path.exists():
            detail = (result.stderr or result.stdout).decode("utf-8", errors="ignore").strip()
            raise RuntimeError(f"DOC 转换失败: 未生成 DOCX 文件。{detail}")

        return _convert_docx_with_python_docx(output_path)


def _convert_mhtml_to_markdown(file_path: Path) -> str:
    """解析 MHTML/MHT 网页归档中的主 HTML 或纯文本正文。"""
    message = BytesParser(policy=policy.default).parsebytes(file_path.read_bytes())
    body = message.get_body(preferencelist=("html", "plain"))
    if body is None:
        raise ValueError("MHTML 文件中未找到可解析的 HTML 或文本正文")

    raw_content = body.get_payload(decode=True)
    if raw_content is None:
        payload = body.get_payload()
        raw_content = payload.encode("utf-8") if isinstance(payload, str) else b""

    charset = body.get_content_charset()
    if not charset and body.get_content_type() == "text/html":
        charset_match = re.search(rb"charset\s*=\s*[\"']?\s*([a-zA-Z0-9._-]+)", raw_content[:4096], re.IGNORECASE)
        if charset_match:
            charset = charset_match.group(1).decode("ascii")

    try:
        content = raw_content.decode(charset or "utf-8", errors="replace")
    except LookupError:
        content = raw_content.decode("utf-8", errors="replace")

    if body.get_content_type() == "text/html":
        return md_convert(content, heading_style="ATX").strip()
    return content.strip()


def _convert_csv_to_markdown(file_path: Path) -> str:
    import pandas as pd

    dataframe = pd.read_csv(file_path)
    tables: list[str] = []
    for i in range(len(dataframe)):
        row_dataframe = dataframe.iloc[[i]]
        tables.append(row_dataframe.to_markdown(index=False))
    return "\n\n".join(tables)


def pdfreader(file_path, params=None):
    """读取 PDF 文件并返回 text 文本。"""
    if isinstance(file_path, str):
        file_path = Path(file_path)

    assert file_path.exists(), "File not found"
    assert file_path.suffix.lower() == ".pdf", "File format not supported"

    loader = PyPDFLoader(str(file_path))
    docs = loader.load()
    text = "\n\n".join([d.page_content for d in docs])
    return text


def parse_pdf(file, params=None):
    """解析 PDF 文件，支持多种 OCR 方式。"""
    from openzetc.knowledge.parser.base import DocumentProcessorException
    from openzetc.knowledge.parser.factory import DocumentProcessorFactory

    opt_ocr, processor_params = _resolve_ocr_engine_params(params)

    if opt_ocr == "disable":
        return pdfreader(file, params=processor_params)

    image_bucket, image_prefix = _resolve_image_storage_params(processor_params)
    processor_params.setdefault("image_bucket", image_bucket)
    processor_params.setdefault("image_prefix", image_prefix)

    try:
        return DocumentProcessorFactory.process_file(opt_ocr, file, processor_params)
    except DocumentProcessorException as e:
        logger.error(f"文档处理失败: {e.service_name} - {str(e)}")
        raise
    except Exception as e:  # noqa: BLE001
        logger.error(f"PDF 解析失败: {str(e)}")
        raise DocumentProcessorException(f"PDF解析失败: {str(e)}", opt_ocr, "parsing_failed")


def parse_image(file, params=None):
    """解析图像文件，支持多种 OCR 方式。"""
    from openzetc.knowledge.parser.base import DocumentProcessorException
    from openzetc.knowledge.parser.factory import DocumentProcessorFactory

    opt_ocr, processor_params = _resolve_ocr_engine_params(params)

    if opt_ocr == "disable":
        raise ValueError(
            "图像文件必须启用OCR才能提取文本内容。"
            "请选择OCR方式 "
            "(rapid_ocr/mineru_ocr/mineru_official/pp_structure_v3_ocr/deepseek_ocr/"
            "paddleocr_vl_1_6/paddleocr_pp_ocrv6) 或移除该文件。"
        )

    image_bucket, image_prefix = _resolve_image_storage_params(processor_params)
    processor_params.setdefault("image_bucket", image_bucket)
    processor_params.setdefault("image_prefix", image_prefix)

    try:
        return DocumentProcessorFactory.process_file(opt_ocr, file, processor_params)
    except DocumentProcessorException as e:
        logger.error(f"图像处理失败: {e.service_name} - {str(e)}")
        raise
    except Exception as e:  # noqa: BLE001
        logger.error(f"图像解析失败: {str(e)}")
        raise DocumentProcessorException(f"图像解析失败: {str(e)}", opt_ocr, "parsing_failed")


async def parse_pdf_async(file, params=None):
    return await asyncio.to_thread(parse_pdf, file, params=params)


async def parse_image_async(file, params=None):
    return await asyncio.to_thread(parse_image, file, params=params)


async def _process_file_to_markdown_core(
    file_path: str, params: dict | None = None
) -> tuple[str, str | None, dict[str, Any]]:
    """将不同类型的文件转换为 markdown，支持本地文件和 MinIO 文件。"""
    from openzetc.knowledge.utils.kb_utils import is_minio_url, parse_minio_url
    from openzetc.storage.minio.client import get_minio_client

    if is_minio_url(file_path):
        logger.debug(f"Downloading file from MinIO: {file_path}")

        if "?" in file_path:
            file_path_clean = file_path.split("?")[0]
        else:
            file_path_clean = file_path

        original_filename = file_path_clean.split("/")[-1]

        with tempfile.NamedTemporaryFile(delete=False, suffix=Path(original_filename).suffix) as temp_file:
            temp_path = temp_file.name

        try:
            bucket_name, object_name = parse_minio_url(file_path)
            minio_client = get_minio_client()
            file_content = await minio_client.adownload_file(bucket_name, object_name)

            async with aiofiles.open(temp_path, "wb") as f:
                await f.write(file_content)

            logger.debug(f"File downloaded to temp path: {temp_path}")
            actual_file_path = temp_path

        except Exception as e:  # noqa: BLE001
            if os.path.exists(temp_path):
                os.unlink(temp_path)
            logger.error(f"Failed to download file from MinIO: {e}")
            raise ValueError(f"无法从MinIO下载文件: {e}")
    else:
        actual_file_path = file_path

    file_ext: str | None = None
    artifacts: dict[str, Any] = {}

    try:
        file_path_obj = Path(actual_file_path)
        file_ext = file_path_obj.suffix.lower()

        if file_ext == ".pdf":
            text = await parse_pdf_async(str(file_path_obj), params=params)
            result = f"{text}"

        elif file_ext in [".txt", ".md"]:
            async with aiofiles.open(file_path_obj, encoding="utf-8") as f:
                content = await f.read()
            result = f"{content}"

        elif file_ext == ".docx":
            try:
                result = await asyncio.to_thread(_convert_with_docling, file_path_obj, params=params)
            except Exception as e:  # noqa: BLE001
                logger.warning(f"Docling 解析 DOCX 失败，回退到 python-docx: {file_path_obj.name}, {e}")
                result = await asyncio.to_thread(_convert_docx_with_python_docx, file_path_obj)

        elif file_ext == ".pptx":
            result = await asyncio.to_thread(_convert_with_docling, file_path_obj, params=params)

        elif file_ext == ".doc":
            result = await asyncio.to_thread(_convert_doc_with_libreoffice, file_path_obj)

        elif file_ext in [".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".tif"]:
            text = await parse_image_async(str(file_path_obj), params=params)
            result = f"{text}"

        elif file_ext in [".html", ".htm"]:
            async with aiofiles.open(file_path_obj, encoding="utf-8") as f:
                content = await f.read()
            text = await asyncio.to_thread(md_convert, content, heading_style="ATX")
            result = f"{text}"

        elif file_ext in [".mhtml", ".mht"]:
            result = await asyncio.to_thread(_convert_mhtml_to_markdown, file_path_obj)

        elif file_ext == ".csv":
            result = await asyncio.to_thread(_convert_csv_to_markdown, file_path_obj)

        elif file_ext in [".xls", ".xlsx"]:
            result = await asyncio.to_thread(_convert_with_docling, file_path_obj, params=params)

        elif file_ext == ".json":
            import json

            async with aiofiles.open(file_path_obj, encoding="utf-8") as f:
                content = await f.read()
            data = json.loads(content)
            json_str = json.dumps(data, ensure_ascii=False, indent=2)
            result = f"```json\n{json_str}\n```"

        elif file_ext == ".zip":
            image_bucket, image_prefix = _resolve_image_storage_params(params)
            zip_result = await _process_zip_file(
                str(file_path_obj),
                image_bucket=image_bucket,
                image_prefix=image_prefix,
            )

            artifacts = {
                "zip_images_info": zip_result["images_info"],
                "zip_content_hash": zip_result["content_hash"],
                "zip_image_bucket": image_bucket,
                "zip_image_prefix": image_prefix,
            }

            result = zip_result["markdown_content"]

        else:
            raise ValueError(f"Unsupported file type: {file_ext}")

    except Exception:
        if is_minio_url(file_path) and os.path.exists(actual_file_path):
            try:
                os.unlink(actual_file_path)
                logger.debug(f"Cleaned up temp file: {actual_file_path}")
            except Exception as cleanup_e:  # noqa: BLE001
                logger.warning(f"Failed to clean up temp file {actual_file_path}: {cleanup_e}")
        raise

    finally:
        if is_minio_url(file_path) and os.path.exists(actual_file_path):
            try:
                os.unlink(actual_file_path)
                logger.debug(f"Cleaned up temp file: {actual_file_path}")
            except Exception as e:  # noqa: BLE001
                logger.warning(f"Failed to clean up temp file {actual_file_path}: {e}")

    return result, file_ext, artifacts


async def parse_source_to_markdown(source: str, params: dict | None = None) -> MarkdownParseResult:
    """统一入口: 将文件解析为 Markdown（URL 解析已废弃）。"""
    markdown, file_ext, artifacts = await _process_file_to_markdown_core(source, params=params)
    return MarkdownParseResult(
        markdown=markdown,
        file_ext=file_ext,
        artifacts=artifacts,
    )


class Parser:
    """Lightweight facade for converting file sources to markdown."""

    @staticmethod
    async def aparse(source: str, params: dict | None = None) -> str:
        """Asynchronously parse source content and return markdown text."""
        parsed = await parse_source_to_markdown(source=source, params=params)
        return parsed.markdown

    @classmethod
    def parse(cls, source: str, params: dict | None = None) -> str:
        """Synchronously parse source content and return markdown text."""
        try:
            asyncio.get_running_loop()
        except RuntimeError:
            return asyncio.run(cls.aparse(source=source, params=params))

        raise RuntimeError("当前处于异步上下文，请使用 `await Parser.aparse(...)`")
